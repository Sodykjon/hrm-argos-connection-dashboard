"""
Takroriy STIR: bir STIR ostidagi reestr satrlaridan qaysi biri ARGOS daraxtidagi tugun ekanini topadi.

Foydalanuvchi qoidasi (07.10.2026): bir STIR bir nechta muassasada bo'lsa — daraxtda bor muassasa
ulangan, daraxtda yo'qlari ulanmagan. Daraxt nomlari ruscha, reestr nomlari o'zbekcha, shuning uchun
moslik bir marta shu skript bilan topiladi va data/argos-stir-match.json ga yoziladi; jonli hisob
(lib/argos-live.ts) faqat shu fayldagi daraxt nomini qidiradi.

Kirish:  --tree argos_tree_YYYY-MM-DD.json (tin, billing, label, path)   --base data/argos-live-base.json
Chiqish: --out data/argos-stir-match.json   --review <xlsx> (noaniq juftliklar, qo'lda ko'rish uchun)
         --keep  oldingi match faylidagi «manual» yozuvlar saqlanadi

Tuzilma solishtiruvi (hudud, tuman skeleti, tur, raqam, joy nomi) scripts/compare-argos-tree.py dan olinadi.
"""
import argparse, collections, json, os, re, sys

HERE = os.path.dirname(os.path.abspath(__file__))
_src = open(os.path.join(HERE, "compare-argos-tree.py"), encoding="utf-8").read().split("\nap = argparse")[0]
H = {}
exec(compile(_src, "compare-argos-tree.py", "exec"), H)
kind, numbers, places, tokens, share, city = H["kind"], H["numbers"], H["places"], H["tokens"], H["share"], H["city"]
tree_district, district_skel, tree_region, base_region = H["tree_district"], H["district_skel"], H["tree_region"], H["base_region"]

ap = argparse.ArgumentParser()
ap.add_argument("--tree", required=True)
ap.add_argument("--base", default="data/argos-live-base.json")
ap.add_argument("--out", default="data/argos-stir-match.json")
ap.add_argument("--review")
ap.add_argument("--keep", action="store_true")
a = ap.parse_args()
sys.stdout.reconfigure(encoding="utf-8")

T = json.load(open(a.tree, encoding="utf-8"))
B = json.load(open(a.base, encoding="utf-8"))
nodes = collections.defaultdict(list)
for r in T["rows"]:
    nodes[str(r["tin"]).strip()].append(r)
groups = collections.defaultdict(list)
for o in B["orgs"]:
    if o.get("stir"):
        groups[o["stir"]].append(o)
dups = {s: l for s, l in groups.items() if len(l) > 1}

old = {}
if a.keep and os.path.exists(a.out):
    for m in json.load(open(a.out, encoding="utf-8"))["matches"]:
        if m.get("how") == "manual":
            old[(m["stir"], m["region"], m["name"])] = m


def score(r, o):
    """(ok, points, reasons) — ok=False means a hard contradiction."""
    why, pts = [], 0
    tr, br = tree_region(r), base_region(o)
    if tr and br and tr != br:
        return False, 0, ["region"]
    kt, kb = kind(r["label"]), kind(o["name"])
    if kt == kb:
        pts += 3
        why.append("type")
    elif kt != "other" and kb != "other":
        return False, 0, ["type %s/%s" % (kt, kb)]
    td = tree_district(r)
    bd = district_skel(o["district"] or "") if o["region"] != "Республика муассасалари" else district_skel(o["name"])
    if td and bd:
        if td != bd:
            pts -= 2
            why.append("district?")
        else:
            pts += 1
            why.append("district")
    nt, nb = numbers(r["label"]), numbers(o["name"])
    if nt and nb:
        if nt & nb:
            pts += 3
            why.append("number")
        else:
            return False, 0, ["number"]
    elif nt or nb:
        pts -= 1
    pt, pb = places(r["label"]) | tokens(r["label"]), places(o["name"]) | tokens(o["name"])
    if share(pt, pb, {td, bd}):
        pts += 2
        why.append("place")
    ct, cb = city(r["label"]), city(o["name"])
    if ct is not None and cb is not None and ct != cb:
        pts -= 2
        why.append("city?")
    if re.search(r"filial|филиал", o["name"], re.I) and not re.search(r"филиал", r["label"], re.I):
        pts -= 2  # a branch is rarely the tree node of its parent's STIR
        why.append("branch")
    return True, pts, why


out, review = [], []
for stir, rows in dups.items():
    for n in nodes.get(stir, []):
        kept = [m for k, m in old.items() if k[0] == stir and m["label"] == n["label"]]
        if kept:
            out.extend(kept)
            continue
        cand = []
        for o in rows:
            ok, pts, why = score(n, o)
            if ok:
                cand.append((pts, o, why))
        cand.sort(key=lambda x: -x[0])
        if not cand:
            review.append({"stir": stir, "label": n["label"], "billing": n["billing"], "pick": None, "why": "nomzod yo'q", "rows": rows})
            continue
        best = cand[0]
        sure = best[0] >= 3 and (len(cand) == 1 or best[0] - cand[1][0] >= 2)
        rec = {"stir": stir, "region": best[1]["region"], "district": best[1].get("district"), "name": best[1]["name"],
               "label": n["label"], "how": "auto" if sure else "review", "why": best[2], "pts": best[0]}
        out.append(rec)
        if not sure:
            review.append({"stir": stir, "label": n["label"], "billing": n["billing"], "pick": best[1]["name"],
                           "why": ", ".join(best[2]) + f" ({best[0]} / keyingi {cand[1][0] if len(cand) > 1 else '-'})",
                           "rows": rows})

json.dump({"source": os.path.basename(a.tree), "at": T.get("at"),
           "rule": "bir STIR ostida: daraxtdagi tugunga mos satr — billing bo'yicha, qolganlari — ulanmagan",
           "matches": out}, open(a.out, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
c = collections.Counter(m["how"] for m in out)
print(json.dumps({"dup_stirs": len(dups), "dup_rows": sum(len(v) for v in dups.values()),
                  "stirs_in_tree": sum(1 for s in dups if s in nodes), "matched": len(out), **c,
                  "no_candidate": sum(1 for r in review if r["pick"] is None)}, ensure_ascii=False))

if a.review:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Takroriy STIR moslik"
    ws.append(["STIR", "ARGOS daraxtidagi nom", "Billing", "Tanlangan reestr satri", "Sabab", "Shu STIR ostidagi barcha reestr satrlari", "Qaror (to'g'ri / boshqa satr nomi / hech biri)"])
    for c_ in ws[1]:
        c_.font = Font(bold=True, color="FFFFFF")
        c_.fill = PatternFill("solid", fgColor="305496")
        c_.alignment = Alignment(wrap_text=True, vertical="center")
    for r in review:
        ws.append([r["stir"], r["label"], "faol" if r["billing"] else "nofaol", r["pick"] or "— topilmadi —", r["why"],
                   "\n".join(f"{o['region']} › {o.get('district') or ''} › {o['name']}" for o in r["rows"]), ""])
        for c_ in ws[ws.max_row]:
            c_.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDEFG", (12, 50, 8, 50, 30, 80, 30)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    wb.save(a.review)
    print("review:", len(review), "->", a.review)

"""
/argos-jonli sahifasi uchun bazaviy reestr ro'yxatini quradi.

Kirish:
  --orig  HRM_ARGOS_17.09.2026_BARCHA_HUDUDLAR_1 (4).xlsx        (17.09 reestr holati)
  --cur   HRM_ARGOS_17.09.2026_BARCHA_HUDUDLAR_1 (4)_yangilangan.xlsx  (joriy ro'yxat: qaysi satrlar qolgan)
  --overrides overrides.json   ({"faol": [[varaq, nom, stir], ...], "ulanmagan": [...], "faolStir": [stir, ...]})
Chiqish:
  data/argos-live-base.json
  (ixtiyoriy) --expected tests/fixtures/argos-live-expected.json — joriy Excel holatlari, test uchun

Satr kaliti: varaq + nom + STIR (satr raqami emas — foydalanuvchi satr o'chiradi/qo'shadi).
"""
import argparse, json
from collections import defaultdict
import openpyxl

REGION = {
    "Respublika markazlari": "Республика муассасалари",
    "Qoraqalpogʻiston": "Қорақалпоғистон Республикаси",
    "Andijon": "Андижон вилояти", "Buxoro": "Бухоро вилояти", "Jizzax": "Жиззах вилояти",
    "Qashqadaryo": "Қашқадарё вилояти", "Navoiy": "Навоий вилояти", "Namangan": "Наманган вилояти",
    "Samarqand": "Самарқанд вилояти", "Surxondaryo": "Сурхондарё вилояти", "Sirdaryo": "Сирдарё вилояти",
    "Toshkent viloyati": "Тошкент вилояти", "Fargʻona": "Фарғона вилояти", "Xorazm": "Хоразм вилояти",
    "Toshkent shahri": "Тошкент шаҳри",
}
ST = {"Ulangan (faol)": "ulangan", "Ulanmagan": "ulanmagan", "Tizimdan oʻchirilgan": "ochirilgan"}

def cols(sn):  # nom, STIR, holat
    return (2, 3, 4) if sn == "Respublika markazlari" else (3, 4, 5)

def contract_col(sn):  # «Shartnoma» ustuni
    return 5 if sn == "Respublika markazlari" else 6

def rows(ws):
    for r in range(10, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if isinstance(a, int) or (isinstance(a, str) and a.strip().isdigit()):
            yield r

def key(ws, sn, r):
    n, s, _ = cols(sn)
    return (sn, str(ws.cell(r, n).value or "").strip(), str(ws.cell(r, s).value or "").strip())

ap = argparse.ArgumentParser()
ap.add_argument("--orig", required=True); ap.add_argument("--cur", required=True)
ap.add_argument("--overrides", required=True); ap.add_argument("--out", default="data/argos-live-base.json")
ap.add_argument("--expected")
a = ap.parse_args()

ov = json.load(open(a.overrides, encoding="utf-8"))
OV_FA = {tuple(k) for k in ov["faol"]}; OV_UL = {tuple(k) for k in ov["ulanmagan"]}; OV_STIR = set(ov["faolStir"])

orig = openpyxl.load_workbook(a.orig)
reestr = defaultdict(list)
for sn in orig.sheetnames[1:]:
    w = orig[sn]
    for r in rows(w):
        reestr[key(w, sn, r)].append(ST[w.cell(r, cols(sn)[2]).value])

cur = openpyxl.load_workbook(a.cur)
out, expected, used = [], [], defaultdict(int)
for sn in cur.sheetnames[1:]:
    w = cur[sn]
    for r in rows(w):
        k = key(w, sn, r)
        i = used[k]; used[k] += 1
        rs = reestr[k][min(i, len(reestr[k]) - 1)] if reestr[k] else None
        assert rs, f"reestrda topilmadi: {k}"
        o = "ulanmagan" if k in OV_UL else "ulangan" if (k in OV_FA or k[2] in OV_STIR) else None
        out.append({
            "region": REGION[sn],
            "district": None if sn == "Respublika markazlari" else str(w.cell(r, 2).value or "").strip(),
            "name": k[1], "stir": k[2] or None, "reestr": rs, "override": o,
            "contract": (str(w.cell(r, contract_col(sn)).value or "").strip() or None),
        })
        expected.append(ST[w.cell(r, cols(sn)[2]).value])

json.dump({"source": "HRM_ARGOS_17.09.2026 reestr", "orgs": out}, open(a.out, "w", encoding="utf-8"), ensure_ascii=False, separators=(",", ":"))
print("orgs", len(out), "overrides", sum(1 for x in out if x["override"]))
if a.expected:
    json.dump(expected, open(a.expected, "w", encoding="utf-8"), separators=(",", ":"))
